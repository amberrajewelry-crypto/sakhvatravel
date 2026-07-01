#!/usr/bin/env python3
"""Семантическое ядро + ссылки русских страниц sakhva-travel.com → Excel.
Только кириллица, только RU-страницы (без /en/ и служебных). Сортировка по убыванию частоты.
"""
from __future__ import annotations
import re
import subprocess
from collections import Counter
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path("/Users/vladimir/sakhva-travel")

# --- стоп-слова RU (предлоги, союзы, частицы, местоимения, общие глаголы) ---
STOP = set("""
и в во не что он на я с со как а то все она так его но да ты к у же вы за бы по только ее мне
было вот от меня еще нет о из ему теперь когда даже ну вдруг ли если уже или ни быть был него до
вас нибудь опять уж вам ведь там потом себя ничего ей может они тут где есть надо ней для мы тебя
их чем была сам чтоб без будто чего раз тоже себе под будет ж тогда кто этот того потому этого
какой совсем ним здесь этом один почти мой тем чтобы нее сейчас были куда зачем всех никогда можно
при наконец два об другой хоть после над больше тот через эти нас про всего них какая много разве
три эту моя впрочем хорошо свою этой перед иногда лучше чуть том нельзя такой им более всегда
конечно всю между это эта эти весь все наш ваш мой твой свой который которые которых этих этими
очень также можно нужно есть нету будут было были также том той тех такие такая такое таких этому
этих всем всей всеми каждый каждая каждое любой любая нашем нашей нашего вашем вашей около либо
""".split())

# дополнительный мусор: служебные/верстка-слова, которые не часть ядра
JUNK = set("""
javascript функция true false var const let return null undefined https http www com div span
class style script href src img png jpg webp svg href rel nofollow noopener target blank
читать подробнее далее меню закрыть открыть назад вперед страница главная cookie cookies
""".split())


def is_ru_page(p: Path) -> bool:
    s = str(p)
    if "/en/" in s or "/node_modules/" in s:
        return False
    bad = ["_bak", "backup", "preview", "demo", "proto", "hero3d",
           "dashboard", "/partner/", "miralinks"]
    if any(b in s for b in bad):
        return False
    name = p.name
    if name == "404.html":
        return False
    if re.match(r"google[0-9a-f]+\.html", name) or name.startswith("yandex_") \
       or re.match(r"sakhvatravel.*\.html", name):
        return False
    return True


def rel_url(p: Path) -> str:
    r = "/" + str(p.relative_to(ROOT))
    r = r.replace("/index.html", "/")
    return r


CYR = re.compile(r"[а-яё]+", re.IGNORECASE)


def tokens(text: str) -> list[str]:
    out = []
    for w in CYR.findall(text.lower()):
        if len(w) < 3:
            continue
        if w in STOP or w in JUNK:
            continue
        out.append(w)
    return out


def main() -> None:
    files = sorted(p for p in ROOT.rglob("*.html") if is_ru_page(p))
    print(f"RU-страниц: {len(files)}")

    word_freq = Counter()          # частота слов в контенте (всего вхождений)
    bigram_freq = Counter()
    trigram_freq = Counter()
    meta_kw_docfreq = Counter()    # сколько страниц используют ключ в meta keywords
    title_kw = Counter()           # слова из title/h1 (вес ядра)
    int_links = Counter()          # внутренние ссылки (RU)
    int_anchor: dict[str, str] = {}
    ext_links = Counter()
    pages_scanned = 0

    for p in files:
        try:
            html = p.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            continue
        soup = BeautifulSoup(html, "html.parser")

        # meta keywords (document frequency — настоящее ядро)
        mk = soup.find("meta", attrs={"name": "keywords"})
        if mk and mk.get("content"):
            seen = set()
            for kw in mk["content"].split(","):
                kw = kw.strip().lower()
                if kw and CYR.search(kw) and kw not in seen:
                    meta_kw_docfreq[kw] += 1
                    seen.add(kw)

        # title + h1/h2 → усиленный вес слов ядра
        for sel in ["title", "h1", "h2"]:
            for el in soup.find_all(sel):
                for w in tokens(el.get_text(" ")):
                    title_kw[w] += 1

        # ссылки
        for a in soup.find_all("a", href=True):
            href = a["href"].strip()
            if not href or href.startswith(("#", "tel:", "mailto:", "javascript:")):
                continue
            anchor = a.get_text(" ", strip=True)[:80]
            low = href.lower()
            is_ext = low.startswith("http") and "sakhva-travel.com" not in low
            if is_ext:
                ext_links[href] += 1
            else:
                # внутренняя; нормализуем, исключаем англ
                path = re.sub(r"^https?://sakhva-travel\.com", "", href)
                if path.startswith("/en/") or path == "/en":
                    continue
                path = path.split("#")[0].split("?")[0]
                if not path:
                    path = "/"
                int_links[path] += 1
                if path not in int_anchor and anchor:
                    int_anchor[path] = anchor

        # видимый контент: убрать script/style/noscript + сквозную обвязку
        for t in soup(["script", "style", "noscript", "nav", "header", "footer"]):
            t.decompose()
        chrome = re.compile(
            r"(nav|menu|footer|header|cookie|fab|modal|breadcrumb|sticky|"
            r"drawer|burger|topbar|hidden|sr-only|skip)", re.I)
        for el in soup.find_all(attrs={"class": chrome}):
            el.decompose()
        for el in soup.find_all(attrs={"id": chrome}):
            el.decompose()
        text = soup.get_text(" ")
        toks = tokens(text)
        word_freq.update(toks)
        for i in range(len(toks) - 1):
            bigram_freq[f"{toks[i]} {toks[i+1]}"] += 1
        for i in range(len(toks) - 2):
            trigram_freq[f"{toks[i]} {toks[i+1]} {toks[i+2]}"] += 1
        pages_scanned += 1

    print(f"Просканировано: {pages_scanned}")
    print(f"meta-ключей: {len(meta_kw_docfreq)}, слов: {len(word_freq)}, "
          f"внутр.ссылок: {len(int_links)}, внеш.ссылок: {len(ext_links)}")

    # --- Excel ---
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1A3D2E")
    head_font = Font(color="FFFFFF", bold=True, size=11)

    def sheet(title: str, headers: list[str], rows: list[tuple], widths: list[int]):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(1, c)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        return ws

    wb.remove(wb.active)

    # 1. meta keywords — ядро по числу страниц
    sheet("Ядро (meta keywords)",
          ["Ключевой запрос", "Страниц использует", "Слов в запросе"],
          [(k, n, len(k.split())) for k, n in meta_kw_docfreq.most_common()],
          [55, 18, 14])

    # 2. фразы 2 слова
    sheet("Фразы 2 слова",
          ["Фраза", "Частота (вхождений)"],
          [(k, n) for k, n in bigram_freq.most_common() if n >= 3],
          [50, 20])

    # 3. фразы 3 слова
    sheet("Фразы 3 слова",
          ["Фраза", "Частота (вхождений)"],
          [(k, n) for k, n in trigram_freq.most_common() if n >= 3],
          [55, 20])

    # 4. отдельные слова
    sheet("Слова (контент)",
          ["Слово", "Частота (вхождений)"],
          [(k, n) for k, n in word_freq.most_common() if n >= 3],
          [35, 20])

    # 5. слова из title/h1/h2 — приоритет ядра
    sheet("Слова в Title-H1-H2",
          ["Слово", "Частота в заголовках"],
          [(k, n) for k, n in title_kw.most_common()],
          [35, 22])

    # 6. внутренние ссылки
    sheet("Внутренние ссылки (RU)",
          ["URL", "Вхождений по сайту", "Пример анкора"],
          [(u, n, int_anchor.get(u, "")) for u, n in int_links.most_common()],
          [55, 20, 45])

    # 7. внешние ссылки
    sheet("Внешние ссылки",
          ["URL", "Вхождений по сайту"],
          [(u, n) for u, n in ext_links.most_common()],
          [70, 20])

    out = ROOT / "scripts" / "semcore_ru.xlsx"
    wb.save(out)
    print(f"Сохранено: {out}")
    # сводка для консоли
    print("\nTOP-15 ядро (meta keywords):")
    for k, n in meta_kw_docfreq.most_common(15):
        print(f"  {n:>3}  {k}")


if __name__ == "__main__":
    main()
