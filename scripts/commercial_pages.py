#!/usr/bin/env python3
"""Полный инвентарь КОММЕРЧЕСКИХ страниц sakhva-travel.com (туры/экскурсии/услуги):
URL + название (H1) + главный ключ (title/meta). Иерархия сверху вниз. → Excel.
"""
from __future__ import annotations
import re
from collections import defaultdict
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path("/Users/vladimir/sakhva-travel")
SITE = "https://sakhva-travel.com"

# Разделы-каталоги, относящиеся к коммерции (туры/экскурсии/услуги/направления)
COMM_DIRS = {
    "ekskursiya", "ekskursii-po-gruzii", "ekskursii-tbilisi", "ekskursii-v-gruzii",
    "ekskursii-v-tbilisi", "ekskursionnye-tury-v-gruziyu",
    "odnodnevnye-ekskursii-iz-tbilisi", "tours", "tour", "tury", "tury-v-gruziyu",
    "tury-v-batumi", "tury-na-kazbek", "tury-v-svaneti", "tury-v-tbilisi",
    "avtorskie-tury-gruzia", "gornye-tury-gruzia", "mnogodnevnye-tury-gruzia",
    "grupovye-tury-v-gruziyu", "individualnyy-tur-v-gruziyu", "vinnye-tury-gruzia",
    "chastniy-gid-tbilisi", "gid-v-tbilisi-na-russkom", "transfer", "prices",
    # регионы Грузии (тур-страницы по регионам)
    "adjara", "guria", "imereti", "kakheti", "kvemo-kartli", "mtskheta-mtianeti",
    "racha", "samegrelo", "samtskhe-javakheti", "shida-kartli",
}
# плюс любые tury-v-gruziyu-iz-* / -letom / -zimoy / -s-detmi
COMM_PREFIX = ("tury-v-gruziyu-",)

EXCLUDE_DIRS = {"en", "node_modules", "images", "api", "scripts", "css", "js",
                "fonts", "blog", "about", "contacts", "privacy", "terms", "policy",
                "gallery", "music", "data", "docs", "schema", "research",
                "seo-audit", "demo", "dashboard", "partner", "netlify", "ads",
                "yandex-verify"}


def top_dir(p: Path) -> str:
    rel = p.relative_to(ROOT)
    return rel.parts[0] if len(rel.parts) > 1 else ""


def is_commercial_page(p: Path) -> bool:
    s = str(p)
    if any(b in s for b in ("_bak", "backup", "preview", "proto", "hero3d",
                            "parasite-seo", "design-", "qa-", "screenshot",
                            "test-results", "miralinks")):
        return False
    d = top_dir(p)
    if not d or d in EXCLUDE_DIRS:
        return False
    if d in COMM_DIRS or d.startswith(COMM_PREFIX):
        return True
    return False


def canon(soup: BeautifulSoup, p: Path) -> str:
    c = soup.find("link", rel="canonical")
    if c and c.get("href"):
        return c["href"].strip()
    return SITE + "/" + str(p.relative_to(ROOT)).replace("/index.html", "/")


# Кластеры верхнего уровня (приоритет сверху вниз) — по пути и H1
def cluster_of(path: str, h1: str, d: str) -> str:
    t = (path + " " + h1).lower()
    if d.startswith("tury-v-gruziyu-iz") or re.search(r"из (москв|спб|санкт|нальчик|"
            r"владикавказ|пятигорск|кисловодск|краснодар|ростов|сочи|воронеж|самар|"
            r"казан|пермь|перми|тюмен|челябинск|екатеринбург|новосибирск|минск|"
            r"ереван|казахстан|ташкент)", t):
        return "Гео-туры в Грузию из городов РФ/СНГ"
    if "transfer" in t or "трансфер" in t or "аэропорт" in t:
        return "Трансфер"
    if "gid" in d or "гид" in h1.lower():
        return "Частный гид"
    if re.search(r"vinn|винн|degustats|дегустац|вино|чача|chacha", t):
        return "Винные туры и дегустации"
    if re.search(r"gastronom|гастроном|khachapuri|хачапур|hinkali|хинкал|кулинар|ужин", t):
        return "Гастрономические туры / мастер-классы"
    if re.search(r"\d+\s*(дн|day)|многодн|mnogodn|vsya-gruziya|вся грузия|неделя|"
                 r"5-dney|10-dney|2-dnya", t):
        return "Многодневные туры по Грузии"
    if re.search(r"kazbeg|казбег|gudaur|гудаур|gergeti|гергет|ananuri|ананури|"
                 r"stepancminda|степанцминд", t):
        return "Казбеги / Гудаури (1 день)"
    if re.search(r"kakheti|кахет|sighnaghi|сигнаг|сигнах|signagi|alazani|алазан|"
                 r"telavi|телав|bodbe|бодбе", t):
        return "Кахетия / Сигнахи (1 день)"
    if re.search(r"mtskheta|мцхет|jvari|джвар|svetic|светицховел", t):
        return "Мцхета (1 день)"
    if re.search(r"kutaisi|кутаис|prometheus|прометея|okace|окаце|martvili|мартвил|"
                 r"kanyon|каньон|kinchkha|кинчх", t):
        return "Кутаиси / каньоны (1 день)"
    if re.search(r"\bgori\b|гори|uplistsikhe|уплисцих", t):
        return "Гори / Уплисцихе (1 день)"
    if re.search(r"borjomi|боржом|bakuriani|бакуриан", t):
        return "Боржоми / Бакуриани"
    if re.search(r"svaneti|сванет|mestia|мести|ushguli|ушгул", t):
        return "Сванетия"
    if re.search(r"vardzia|вардзи|akhaltsikhe|ахалцих|rabati|рабат|david-gareji|"
                 r"гареджи|gareji", t):
        return "Вардзиа / Давид-Гареджи"
    if re.search(r"batumi|батуми|adjara|аджар|gomis|гомис", t):
        return "Батуми / Аджария"
    if re.search(r"tusheti|тушети|kazbegi-kakheti", t):
        return "Горные / тематические туры"
    if re.search(r"nochn|ночн", t):
        return "Ночные экскурсии по Тбилиси"
    if re.search(r"narikala|нарикал|stary-tbilisi|старый тбилиси|sovetskiy|советск|"
                 r"hram|храм|fotoses|фотосес|art-tur|relokant|релокант|pensioner|"
                 r"пенсионер|family|семейн|obzorn|обзорн|peshe|пеш", t):
        return "Тематические экскурсии по Тбилиси"
    if re.search(r"tbilisi|тбилиси", t):
        return "Экскурсии по Тбилиси (общее)"
    if re.search(r"individualn|индивидуальн|avtorsk|авторск|grupp|групп|"
                 r"ekskursionnye-tury|gornye-tury", t):
        return "Туры по Грузии (тип: индивид/групп/авторск)"
    if re.search(r"price|цен|prices", t):
        return "Цены"
    if re.search(r"gruzi|грузи", t):
        return "Туры по Грузии (общее)"
    return "Прочие коммерческие"


def first_kw(soup: BeautifulSoup) -> str:
    mk = soup.find("meta", attrs={"name": "keywords"})
    if mk and mk.get("content"):
        parts = [x.strip() for x in mk["content"].split(",") if x.strip()]
        if parts:
            return parts[0]
    return ""


def main() -> None:
    rows = []  # (cluster, url, h1, key)
    for p in sorted(ROOT.rglob("*.html")):
        if not is_commercial_page(p):
            continue
        html = p.read_text(encoding="utf-8", errors="ignore")
        soup = BeautifulSoup(html, "html.parser")
        url = canon(soup, p)
        h1el = soup.find("h1")
        h1 = re.sub(r"\s+", " ", h1el.get_text(" ").strip()) if h1el else ""
        title = soup.find("title")
        title = re.sub(r"\s+", " ", title.get_text().strip()) if title else ""
        key = first_kw(soup) or title
        d = top_dir(p)
        cl = cluster_of(str(p.relative_to(ROOT)), h1 or title, d)
        rows.append((cl, url, h1 or title, key))

    tree: dict[str, list[tuple]] = defaultdict(list)
    for cl, url, h1, key in rows:
        tree[cl].append((url, h1, key))
    for cl in tree:
        tree[cl].sort(key=lambda r: r[0])
    order = sorted(tree, key=lambda c: -len(tree[c]))

    print(f"коммерческих страниц: {len(rows)}, кластеров: {len(tree)}\n")
    for cl in order:
        print(f"\n■ {cl}  ({len(tree[cl])})")
        for url, h1, key in tree[cl]:
            print(f"    {url.replace(SITE,'')}")
            print(f"        H1: {h1[:70]}")

    # Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Коммерческие страницы"
    ws.append(["Кластер", "URL страницы", "Название (H1)", "Главный ключ (meta/title)"])
    hf = PatternFill("solid", fgColor="1A3D2E")
    ff = Font(color="FFFFFF", bold=True)
    cf = Font(bold=True, color="1A3D2E")
    sub = PatternFill("solid", fgColor="EAF3EE")
    for c in range(1, 5):
        ws.cell(1, c).fill = hf
        ws.cell(1, c).font = ff
    for cl in order:
        rr = ws.max_row + 1
        ws.cell(rr, 1, f"{cl}  ({len(tree[cl])})").font = cf
        for cc in range(1, 5):
            ws.cell(rr, cc).fill = sub
        for url, h1, key in tree[cl]:
            ws.append(["", url, h1, key])
    for i, w in enumerate([44, 58, 55, 45], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    out = ROOT / "scripts" / "commercial_pages.xlsx"
    wb.save(out)
    print(f"\nСохранено: {out}  (страниц: {len(rows)})")


if __name__ == "__main__":
    main()
