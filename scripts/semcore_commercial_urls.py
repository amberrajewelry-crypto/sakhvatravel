#!/usr/bin/env python3
"""Коммерческие ключи → конкретные URL страниц sakhva-travel.com, которые под них
заточены (ключ присутствует в meta keywords страницы). Иерархия сверху вниз.
"""
from __future__ import annotations
import re
from collections import defaultdict
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path("/Users/vladimir/sakhva-travel")
SITE = "https://sakhva-travel.com"


def is_ru_page(p: Path) -> bool:
    s = str(p)
    if "/en/" in s or "/node_modules/" in s:
        return False
    bad = ["_bak", "backup", "preview", "demo", "proto", "hero3d",
           "dashboard", "/partner/", "miralinks"]
    if any(b in s for b in bad):
        return False
    name = p.name
    if name == "404.html":
        return False
    if re.match(r"google[0-9a-f]+\.html", name) or name.startswith("yandex_") \
       or re.match(r"sakhvatravel.*\.html", name):
        return False
    return True


def canon(soup: BeautifulSoup, p: Path) -> str:
    c = soup.find("link", rel="canonical")
    if c and c.get("href"):
        return c["href"].strip()
    return SITE + "/" + str(p.relative_to(ROOT)).replace("/index.html", "/")


COMM_STEMS = ("экскурс", "тур", "гид", "трансфер", "бронир", "заказ", "куп",
              "цен", "стоим", "аренд", "прокат", "такси", "поездк", "путев",
              "путёв", "винн", "дегуст", "индивид", "частн", "однодн", "многодн")
GEO_FROM = re.compile(r"\bиз\s+(тбилиси|батуми|кутаиси|москв|нальчик|владикавказ|"
                      r"минеральн|минвод|ставропол|краснодар|пятигорск|грозн|"
                      r"махачкал|ереван|баку)")
INFO_BLOCK = re.compile(r"(что посмотр|как добра|нужна ли|можно ли|почему|погод|"
                        r"история|достопримеч|карта|виза|сколько денег|бюджет|"
                        r"что попроб|рецепт)")


def toks(s: str) -> list[str]:
    return re.findall(r"[а-яё]+", s.lower())


def is_commercial(phrase: str) -> bool:
    if INFO_BLOCK.search(phrase):
        return False
    if any(t.startswith(COMM_STEMS) for t in toks(phrase)):
        return True
    return bool(GEO_FROM.search(phrase))


CLUSTERS: list[tuple[str, re.Pattern]] = [
    ("Гео-туры из городов РФ/Кавказа",
     re.compile(r"\bиз\s+(москв|нальчик|владикавказ|минеральн|минвод|ставропол|"
                r"краснодар|пятигорск|грозн|махачкал|ереван|баку|кисловодск)")),
    ("Трансфер", re.compile(r"трансфер|аэропорт|такси")),
    ("Частный гид", re.compile(r"\bгид")),
    ("Винные туры (Кахетия)", re.compile(r"винн|дегустац|\bвино\b|квеври|чача")),
    ("Гастрономические туры", re.compile(r"гастроном|кулинарн|ужин|хинкал|хачапур")),
    ("Казбеги / Гудаури", re.compile(r"казбег|степанцминд|гергет|гудаур|ананури|жинвал")),
    ("Кахетия / Сигнахи", re.compile(r"кахет|сигнаг|сигнах|бодбе|телав|цинандал|алаверд")),
    ("Мцхета", re.compile(r"мцхет|джвар|светицховел")),
    ("Кутаиси / каньоны", re.compile(r"кутаис|прометея|окаце|мартвил|каньон|кинчх")),
    ("Гори / Уплисцихе", re.compile(r"\bгори\b|уплисцих|сталин")),
    ("Боржоми / Бакуриани", re.compile(r"боржом|бакуриан")),
    ("Сванетия", re.compile(r"сванет|местиа|ушгул|мести")),
    ("Вардзиа / Самцхе / Давид Гареджи",
     re.compile(r"вардзи|ахалцих|рабат|гареджи|военно")),
    ("Многодневные туры по Грузии",
     re.compile(r"многодн|\d\s*дн|недел|тур по грузии|туры по грузии|"
                r"туры в грузию|тур в грузию")),
    ("Ночной Тбилиси", re.compile(r"ночн")),
    ("Обзорные / пешие по Тбилиси",
     re.compile(r"обзорн|пешеход|\bпеш|прогулк|старый город|серн")),
    ("Экскурсии по Тбилиси (общее)", re.compile(r"тбилиси")),
    ("Цены и бронирование", re.compile(r"цен|стоим|бронир|заказ|предоплат")),
]


def cluster_of(phrase: str) -> str:
    for name, rx in CLUSTERS:
        if rx.search(phrase):
            return name
    return "Прочее коммерческое (аренда/сопутствующее)"


def main() -> None:
    kw_urls: dict[str, list[str]] = defaultdict(list)
    files = sorted(p for p in ROOT.rglob("*.html") if is_ru_page(p))
    for p in files:
        soup = BeautifulSoup(p.read_text(encoding="utf-8", errors="ignore"),
                             "html.parser")
        mk = soup.find("meta", attrs={"name": "keywords"})
        if not (mk and mk.get("content")):
            continue
        url = canon(soup, p)
        for kw in mk["content"].split(","):
            kw = kw.strip().lower()
            if kw and re.search(r"[а-яё]", kw) and is_commercial(kw):
                if url not in kw_urls[kw]:
                    kw_urls[kw].append(url)

    comm = sorted(kw_urls)
    print(f"коммерческих ключей с привязкой: {len(comm)}")

    tree: dict[str, list[str]] = defaultdict(list)
    for k in comm:
        tree[cluster_of(k)].append(k)
    for c in tree:
        tree[c].sort(key=lambda k: (-len(kw_urls[k]), k))
    order = sorted(tree, key=lambda c: (-sum(len(kw_urls[k]) for k in tree[c]),
                                        -len(tree[c])))

    # печать
    for c in order:
        print(f"\n■ {c}")
        for k in tree[c]:
            for i, u in enumerate(kw_urls[k]):
                path = u.replace(SITE, "")
                print(f"    {k:<38} → {path}" if i == 0
                      else f"    {'':<38}   {path}")

    # Excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Ключ → URL"
    ws.append(["Кластер", "Коммерческий ключ", "Страниц", "URL(ы) на сайте"])
    hf = PatternFill("solid", fgColor="1A3D2E")
    ff = Font(color="FFFFFF", bold=True)
    cf = Font(bold=True, color="1A3D2E")
    sub = PatternFill("solid", fgColor="EAF3EE")
    for c in range(1, 5):
        ws.cell(1, c).fill = hf
        ws.cell(1, c).font = ff
    for c in order:
        rr = ws.max_row + 1
        ws.cell(rr, 1, c).font = cf
        for cc in range(1, 5):
            ws.cell(rr, cc).fill = sub
        for k in tree[c]:
            ws.append(["", k, len(kw_urls[k]), "\n".join(kw_urls[k])])
            ws.cell(ws.max_row, 4).alignment = \
                __import__("openpyxl").styles.Alignment(wrap_text=True, vertical="top")
    for i, w in enumerate([34, 40, 9, 70], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    out = ROOT / "scripts" / "semcore_commercial_urls.xlsx"
    wb.save(out)
    print(f"\nСохранено: {out}  (ключей: {len(comm)})")


if __name__ == "__main__":
    main()
